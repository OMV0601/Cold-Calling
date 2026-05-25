"""
Spreadsheet output and console summary.

Writes the final leads table to an .xlsx or .csv file and prints a summary
panel to the terminal using rich.

Public surface:
  build_lead_row(business, report, score, pitch) -> LeadRow
  write_output(rows, config, console)            -> Path
  print_summary(rows, scanned_count, console)    -> None
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

from leadscout.audit.models import AuditReport
from leadscout.config import AppConfig
from leadscout.places import Business
from leadscout.scoring import LeadScore, format_gap_list


# ── Row model ─────────────────────────────────────────────────────────────────

@dataclass
class LeadRow:
    """One row in the output spreadsheet — all fields pre-formatted as strings."""

    name: str
    category: str
    search_area: str
    address: str
    phone: str
    contact_email: str
    contact_url: str
    website_url: str
    lead_track: str
    rating: str
    review_count: str
    listing_photo_count: str
    gaps: str
    pitch: str
    lead_score: float       # kept as float so pandas can sort it


# ── Column layout ─────────────────────────────────────────────────────────────

# (header label, LeadRow attribute, max column width in Excel)
_COLUMNS: list[tuple[str, str, int]] = [
    ("Lead Score",         "lead_score",          12),
    ("Track",              "lead_track",           18),
    ("Business Name",      "name",                 32),
    ("Category",           "category",             16),
    ("City / Area",        "search_area",          18),
    ("Rating",             "rating",               8),
    ("Reviews",            "review_count",         10),
    ("Photos (listing)",   "listing_photo_count",  16),
    ("Phone",              "phone",                16),
    ("Contact Email",      "contact_email",        28),
    ("Website",            "website_url",          36),
    ("Contact Page",       "contact_url",          36),
    ("Address",            "address",              36),
    ("Gaps",               "gaps",                 60),
    ("Pitch Opener",       "pitch",                80),
]

_COL_HEADERS = [c[0] for c in _COLUMNS]
_COL_ATTRS   = [c[1] for c in _COLUMNS]
_COL_WIDTHS  = [c[2] for c in _COLUMNS]


# ── Builder ───────────────────────────────────────────────────────────────────

def build_lead_row(
    business: Business,
    report: Optional[AuditReport],
    score: LeadScore,
    pitch: str,
) -> LeadRow:
    """Assemble a LeadRow from a business record, audit report, score, and pitch."""
    track = "A — needs website" if not business.website_uri else "B — has website"
    return LeadRow(
        name=business.name,
        category=business.category,
        search_area=business.search_area,
        address=business.address,
        phone=business.phone or "",
        contact_email=(report.contact_email or "") if report else "",
        contact_url=(report.contact_url or "") if report else "",
        website_url=business.website_uri or "",
        lead_track=track,
        rating=str(business.rating or ""),
        review_count=str(business.review_count or ""),
        listing_photo_count=str(business.photo_count),
        gaps=format_gap_list(report),
        pitch=pitch,
        lead_score=score.total,
    )


# ── Output writers ────────────────────────────────────────────────────────────

def write_output(
    rows: list[LeadRow],
    config: AppConfig,
    console: Console,
) -> Path:
    """
    Write *rows* to the path in config.output.

    Supports "xlsx" (formatted, frozen header, colour-coded tracks) and "csv".
    Creates parent directories automatically.
    Returns the path that was written.
    """
    path = config.output.path
    path.parent.mkdir(parents=True, exist_ok=True)

    df = _to_dataframe(rows)

    if config.output.format == "xlsx":
        _write_xlsx(df, path)
    else:
        _write_csv(df, path)

    console.print(f"[green]Output written to:[/green] {path}  ({len(rows)} leads)")
    return path


def _to_dataframe(rows: list[LeadRow]) -> pd.DataFrame:
    """Convert LeadRow objects to a pandas DataFrame sorted by lead score."""
    data = [
        {header: getattr(row, attr) for header, attr in zip(_COL_HEADERS, _COL_ATTRS)}
        for row in rows
    ]
    df = pd.DataFrame(data, columns=_COL_HEADERS)
    df["_track_order"] = df["Track"].apply(lambda t: 0 if str(t).startswith("A") else 1)
    df = df.sort_values(["_track_order", "Lead Score"], ascending=[True, False])
    df = df.drop(columns=["_track_order"]).reset_index(drop=True)
    return df


def _write_xlsx(df: pd.DataFrame, path: Path) -> None:
    """Write *df* to an xlsx file with formatting: bold header, track colours, frozen row."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Leads")
        ws = writer.sheets["Leads"]

        # ── Header row ────────────────────────────────────────────────────────
        header_fill = PatternFill("solid", fgColor="2E4057")   # dark navy
        header_font = Font(bold=True, color="FFFFFF", size=10)
        center = Alignment(horizontal="center", vertical="center", wrap_text=False)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
        ws.row_dimensions[1].height = 20

        # ── Freeze header ─────────────────────────────────────────────────────
        ws.freeze_panes = "A2"

        # ── Row fill by track ─────────────────────────────────────────────────
        fill_a = PatternFill("solid", fgColor="FFF9C4")    # soft yellow  → Track A
        fill_b = PatternFill("solid", fgColor="E3F2FD")    # soft blue    → Track B

        track_col_idx = _COL_HEADERS.index("Track")

        for excel_row_idx, df_row_idx in enumerate(range(len(df)), start=2):
            track_val = str(df.iloc[df_row_idx]["Track"])
            fill = fill_a if track_val.startswith("A") else fill_b
            for cell in ws[excel_row_idx]:
                cell.fill = fill
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        # ── Wrap text for long-text columns ───────────────────────────────────
        wrap_cols = {"Gaps", "Pitch Opener", "Address", "Website", "Contact Page"}
        for col_idx, header in enumerate(_COL_HEADERS, start=1):
            if header in wrap_cols:
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)

        # ── Column widths ─────────────────────────────────────────────────────
        for col_idx, (header, max_w) in enumerate(_col_width_pairs(), start=1):
            col_letter = get_column_letter(col_idx)
            # Measure max content width in this column (sample first 100 rows)
            col_data = df.iloc[:100][header].astype(str)
            content_w = max((len(v) for v in col_data), default=10)
            ws.column_dimensions[col_letter].width = min(
                max(content_w + 2, len(header) + 2, 8),
                max_w,
            )

        # ── Score column: 2 decimal places ────────────────────────────────────
        score_col = get_column_letter(_COL_HEADERS.index("Lead Score") + 1)
        for cell in ws[score_col][1:]:   # skip header
            cell.number_format = "0.00"

        # ── Auto-filter on header row ─────────────────────────────────────────
        ws.auto_filter.ref = ws.dimensions


def _write_csv(df: pd.DataFrame, path: Path) -> None:
    """Write *df* to a csv file."""
    df.to_csv(path, index=False, encoding="utf-8-sig")  # utf-8-sig for Excel compat


def _col_width_pairs():
    """Yield (header, max_width) pairs."""
    return zip(_COL_HEADERS, _COL_WIDTHS)


# ── Console summary ───────────────────────────────────────────────────────────

def print_summary(
    rows: list[LeadRow],
    scanned_count: int,
    console: Console,
) -> None:
    """
    Print a rich summary panel:
      - Totals: scanned, qualified per track
      - Top 5 leads by score with their top gap
    """
    track_a = [r for r in rows if r.lead_track.startswith("A")]
    track_b = [r for r in rows if r.lead_track.startswith("B")]
    top5 = sorted(rows, key=lambda r: r.lead_score, reverse=True)[:5]

    # ── Stats panel ───────────────────────────────────────────────────────────
    stats = (
        f"[bold]Businesses scanned :[/bold] {scanned_count}\n"
        f"[bold]Qualified leads    :[/bold] {len(rows)} total\n"
        f"  [yellow]Track A (no website)[/yellow] : {len(track_a)}\n"
        f"  [cyan]Track B (has website)[/cyan] : {len(track_b)}"
    )
    console.print(Panel(stats, title="[bold blue]LeadScout Results[/bold blue]", expand=False))

    if not top5:
        return

    # ── Top 5 table ───────────────────────────────────────────────────────────
    table = Table(
        title="Top 5 leads by score",
        show_header=True,
        header_style="bold white on dark_blue",
        border_style="bright_black",
        show_lines=False,
    )
    table.add_column("#",             style="dim",         width=3,  justify="right")
    table.add_column("Business",      style="bold",        width=30)
    table.add_column("Track",         width=4,             justify="center")
    table.add_column("Score",         style="green bold",  width=7,  justify="right")
    table.add_column("Rating",        style="yellow",      width=7,  justify="center")
    table.add_column("Reviews",       width=8,             justify="right")
    table.add_column("Top gap",       style="dim",         width=28)

    for rank, row in enumerate(top5, start=1):
        top_gap = _first_gap(row.gaps)
        track_display = (
            f"[yellow]A[/yellow]" if row.lead_track.startswith("A")
            else f"[cyan]B[/cyan]"
        )
        table.add_row(
            str(rank),
            row.name[:29],
            track_display,
            f"{row.lead_score:.2f}",
            row.rating,
            row.review_count,
            top_gap,
        )

    console.print(table)


def _first_gap(gaps_str: str) -> str:
    """Extract the first gap name from the pipe-separated gaps string."""
    if not gaps_str:
        return "—"
    first = gaps_str.split("|")[0].strip()
    # Strip severity tag: "[HIGH] no_contact_form" -> "no_contact_form"
    return first.split("] ")[-1] if "]" in first else first
